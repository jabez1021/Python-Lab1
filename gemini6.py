import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.oxml.ns import qn

def set_hangul_font(paragraph, font_name="Malgun Gothic"):
    """한글 폰트 설정"""
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def main():
    chrome_options = Options()
    # 크롤링 과정을 확인하기 위해 브라우저 창을 띄운 상태로 진행합니다.
    
    try:
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=chrome_options)
        
        # 1. 네이버 증권 페이지 접속
        driver.get("https://finance.naver.com/sise/")
        time.sleep(3) # 데이터가 완전히 로딩될 때까지 넉넉히 대기
        
        # 2. 인기 검색 종목 영역의 원본 텍스트 수집
        items = driver.find_elements(By.CSS_SELECTOR, "#popularItemList li")
        
        raw_data_list = []
        for item in items[:10]:
            # .text를 사용하면 해당 엘리먼트 내의 모든 텍스트가 한 줄로 가져와집니다.
            raw_data_list.append(item.text.strip())
        
        print("--- 수집된 원본 데이터 ---")
        for i, raw in enumerate(raw_data_list):
            print(f"{i+1}: {raw}")
        
        driver.quit()

        # 3. template.docx 생성 및 원본 데이터 그대로 저장
        doc = Document()
        
        # 조건에 맞는 {title} 표시 추가
        p_title = doc.add_paragraph("{title}")
        set_hangul_font(p_title)

        doc.add_paragraph("\n[웹사이트에서 가져온 데이터 원본]")
        
        # 가공 없이 그대로 한 줄씩 추가
        for raw in raw_data_list:
            p = doc.add_paragraph(raw)
            set_hangul_font(p)

        doc.save('template.docx')
        print("\n'template.docx' 파일에 원본 데이터가 저장되었습니다. 파일을 확인해 보세요.")

    except PermissionError:
        print("오류: template.docx 파일이 열려 있습니다. 워드를 닫고 실행해 주세요.")
    except Exception as e:
        print(f"오류 발생: {e}")

if __name__ == "__main__":
    main()