import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.oxml.ns import qn

def set_hangul_font(paragraph, font_name="Malgun Gothic"):
    """워드 문서 내 한글 깨짐 방지 설정"""
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def main():
    # --- [단계 1] Selenium을 이용한 데이터 수집 (조건 1, 2, 3) ---
    chrome_options = Options()
    # chrome_options.add_argument("--headless") # 필요 시 설정
    
    try:
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=chrome_options)
        
        driver.get("https://finance.naver.com/sise/")
        time.sleep(2)
        
        # 인기 검색 종목명과 현재가 정보를 수집 (데이터 예시와 유사하게 구성)
        # 네이버 증권의 구조에 맞춰 종목명과 가격 정보를 가져옵니다.
        items = driver.find_elements(By.CSS_SELECTOR, "#popularItemList li")
        
        popular_stocks = []
        for i, item in enumerate(items[:10], 1):
            name = item.find_element(By.TAG_NAME, "a").text
            price = item.find_element(By.TAG_NAME, "span").text
            popular_stocks.append(f"{i}.{name} {price}")
        
        print("수집 데이터:", popular_stocks)
        driver.quit()

        # --- [단계 2] 수집된 데이터를 바탕으로 template.docx 생성 ---
        # 조건에 나온 "전달된 템플릿" 역할을 하기 위해 먼저 생성합니다.
        temp_doc = Document()
        # {title} 표시 포함
        p_title = temp_doc.add_paragraph("{title}")
        set_hangul_font(p_title)
        
        # 수집된 10개 종목을 템플릿 본문에 포함
        temp_doc.add_paragraph("\n[실시간 수집 데이터]")
        for stock in popular_stocks:
            p = temp_doc.add_paragraph(stock)
            set_hangul_font(p)
        
        temp_doc.save('template.docx')
        print("1. 수집된 데이터를 포함한 template.docx 생성 완료.")

        # --- [단계 3] template.docx를 불러와서 편집 (조건 4, 5) ---
        doc = Document('template.docx')
        
        # 조건 4: {title} 부분에 제목 추가
        title_text = "오늘의 인기 검색 종목 10개 입니다."
        for paragraph in doc.paragraphs:
            if '{title}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{title}', title_text)
                set_hangul_font(paragraph)

        # 조건 5: 문서의 하단에 List Number 스타일로 10개의 종목을 추가
        doc.add_paragraph("\n[최종 정리 리스트]")
        for stock in popular_stocks:
            # 스타일이 템플릿에 정의되어 있지 않을 경우를 대비해 예외 처리 혹은 기본 추가
            try:
                p = doc.add_paragraph(stock, style='List Number')
            except:
                p = doc.add_paragraph(f"• {stock}")
            set_hangul_font(p)

        # --- [단계 4] result.docx 저장 (조건 6) ---
        doc.save('result.docx')
        print("2. 최종 결과물인 result.docx 저장 완료.")

    except Exception as e:
        print(f"오류 발생: {e}")

if __name__ == "__main__":
    main()