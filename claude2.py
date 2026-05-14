"""
네이버 증권 인기 검색 종목 수집 → result.docx 저장

실행 전 설치:
    pip install selenium webdriver-manager python-docx

사용법:
    python naver_finance_scraper.py
"""

import time
import sys

# ── 라이브러리 임포트 (예외처리 포함) ─────────────────────────────────────────
try:
    from selenium import webdriver
    from selenium.webdriver.chrome.service import Service
    from selenium.webdriver.chrome.options import Options
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.common.exceptions import (
        WebDriverException,
        TimeoutException,
        NoSuchElementException,
    )
    from webdriver_manager.chrome import ChromeDriverManager
except ImportError as e:
    print(f"[오류] 필수 라이브러리 없음: {e}")
    print("설치: pip install selenium webdriver-manager python-docx")
    sys.exit(1)

try:
    from docx import Document
except ImportError:
    print("[오류] python-docx 없음. 설치: pip install python-docx")
    sys.exit(1)


# ── 1. 크롬 드라이버 생성 ──────────────────────────────────────────────────────
def create_driver() -> webdriver.Chrome:
    """
    Chrome WebDriver 생성.
    - headless 옵션 없음 → 브라우저 화면이 실제로 뜹니다.
    - webdriver-manager 로 ChromeDriver 자동 설치/관리.
    """
    options = Options()
    # ※ headless 미사용 → 크롬 창이 화면에 표시됩니다
    options.add_argument("--start-maximized")
    options.add_argument("--disable-notifications")
    options.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )

    try:
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=options)
        return driver
    except WebDriverException as e:
        raise WebDriverException(f"Chrome 드라이버 실행 실패: {e}") from e


# ── 2 & 3. 네이버 증권 접속 + 인기 검색 종목 수집 ─────────────────────────────
def get_popular_stocks(driver: webdriver.Chrome) -> list:
    """
    https://finance.naver.com/sise/ 의 인기 검색 종목 10개를 반환합니다.
    """
    url = "https://finance.naver.com/sise/"

    try:
        driver.get(url)
        print(f"        → 페이지 접속 완료: {url}")
    except WebDriverException as e:
        raise WebDriverException(f"페이지 접속 실패: {e}") from e

    # 페이지 기본 로딩 대기
    time.sleep(3)

    stocks = []

    # ── 셀렉터 후보 (네이버 증권 페이지 구조 기준) ─────────────────────────
    selector_candidates = [
        (By.CSS_SELECTOR, ".aside_popular .item_name"),
        (By.CSS_SELECTOR, "#popularItemList .item_name"),
        (By.CSS_SELECTOR, ".popular_search .item_name"),
        (By.CSS_SELECTOR, ".aside_popular li a"),
        (By.CSS_SELECTOR, "#popularItemList li a"),
        (By.CSS_SELECTOR, ".aside_popular .tit"),
        # iframe 없이 바로 접근 가능한 경우
        (By.XPATH, '//div[contains(@class,"popular")]//li/a'),
        (By.XPATH, '//ul[contains(@class,"popular")]//li/a'),
    ]

    for by, selector in selector_candidates:
        try:
            WebDriverWait(driver, 5).until(
                EC.presence_of_element_located((by, selector))
            )
            elements = driver.find_elements(by, selector)
            candidates = [el.text.strip() for el in elements if el.text.strip()]
            if len(candidates) >= 5:
                stocks = candidates
                print(f"        → 셀렉터 성공: {selector}  ({len(stocks)}개 수집)")
                break
        except (TimeoutException, NoSuchElementException):
            continue

    # ── iframe 내부 확인 (네이버 증권은 일부 콘텐츠를 iframe에 넣기도 함) ──
    if not stocks:
        try:
            iframes = driver.find_elements(By.TAG_NAME, "iframe")
            for iframe in iframes:
                try:
                    driver.switch_to.frame(iframe)
                    elements = driver.find_elements(By.CSS_SELECTOR, ".item_name, li a")
                    candidates = [el.text.strip() for el in elements if el.text.strip()]
                    if len(candidates) >= 5:
                        stocks = candidates
                        print(f"        → iframe 내부에서 {len(stocks)}개 수집")
                        break
                except Exception:
                    pass
                finally:
                    driver.switch_to.default_content()
        except Exception:
            pass

    if not stocks:
        raise RuntimeError(
            "인기 검색 종목을 찾지 못했습니다.\n"
            "네이버 증권 페이지 구조가 변경됐을 수 있습니다.\n"
            "브라우저 창에서 직접 셀렉터를 확인해 주세요."
        )

    return stocks[:10]


# ── 4 · 5 · 6. Word 문서 작성 및 저장 ────────────────────────────────────────
def save_to_docx(
    stocks: list,
    template_path: str = "template.docx",
    output_path: str = "result.docx",
) -> None:
    """
    template.docx 열기 → {title} 교체 → List Number 종목 추가 → result.docx 저장
    """
    title_text = "오늘의 인기 검색 종목 10개 입니다."

    # 템플릿 열기
    try:
        doc = Document(template_path)
    except Exception as e:
        raise FileNotFoundError(
            f"템플릿 파일을 열 수 없습니다 ({template_path}): {e}"
        ) from e

    # 4. {title} → 제목 교체
    replaced = False
    for para in doc.paragraphs:
        if "{title}" in para.text:
            # run 분리 여부와 관계없이 전체 텍스트 교체
            for run in para.runs:
                run.text = ""                     # 기존 run 초기화
            if para.runs:
                para.runs[0].text = title_text    # 첫 run에 제목 삽입
            else:
                para.add_run(title_text)
            replaced = True
            print(f"        → {{title}} 교체 완료: \"{title_text}\"")
            break

    if not replaced:
        print("[경고] {title} 플레이스홀더를 찾지 못해 문서 상단에 제목을 추가합니다.")
        doc.paragraphs[0].insert_paragraph_before(title_text)

    # 5. List Number 스타일로 종목 10개 추가
    for stock in stocks:
        doc.add_paragraph(stock, style="List Number")
    print(f"        → 종목 {len(stocks)}개를 'List Number' 스타일로 추가")

    # 6. result.docx 저장
    try:
        doc.save(output_path)
        print(f"        → '{output_path}' 저장 완료 ✓")
    except Exception as e:
        raise IOError(f"문서 저장 실패 ({output_path}): {e}") from e


# ── 메인 ──────────────────────────────────────────────────────────────────────
def main():
    driver = None
    try:
        # 1단계: 크롬 실행
        print("[1단계] Chrome 브라우저를 실행합니다 (화면에 창이 표시됩니다)...")
        driver = create_driver()
        print("        → Chrome 실행 성공")

        # 2·3단계: 네이버 증권 접속 + 종목 수집
        print("\n[2·3단계] 네이버 증권 접속 및 인기 검색 종목 수집 중...")
        stocks = get_popular_stocks(driver)

        print("\n  ┌─ 인기 검색 종목 10개 ─────────────────┐")
        for i, s in enumerate(stocks, 1):
            print(f"  │  {i:>2}. {s:<20}          │")
        print("  └───────────────────────────────────────┘")

        # 4·5·6단계: result.docx 생성
        print("\n[4·5·6단계] result.docx 문서 생성 중...")
        save_to_docx(stocks, template_path="template.docx", output_path="result.docx")

        print("\n✅ 완료! result.docx 파일이 현재 폴더에 생성되었습니다.")

    except WebDriverException as e:
        print(f"\n[오류] 브라우저 오류: {e}")
        sys.exit(1)
    except TimeoutException as e:
        print(f"\n[오류] 페이지 로딩 타임아웃: {e}")
        sys.exit(1)
    except (FileNotFoundError, IOError, RuntimeError) as e:
        print(f"\n[오류] {e}")
        sys.exit(1)
    except Exception as e:
        print(f"\n[예상치 못한 오류] {type(e).__name__}: {e}")
        sys.exit(1)
    finally:
        if driver is not None:
            driver.quit()
            print("[정리] Chrome 브라우저를 종료했습니다.")


if __name__ == "__main__":
    main()
