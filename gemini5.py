import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.oxml.ns import qn

def set_hangul_font(paragraph, font_name="Malgun Gothic"):
    """한글 깨짐 방지 폰트 설정"""
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def main():
    chrome_options = Options()
    # chrome_options.add_argument("--headless")
    
    try:
        # --- 1. Selenium 데이터 수집 (등락 정보 포함) ---
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=chrome_options)
        
        driver.get("https://finance.naver.com/sise/")
        time.sleep(2)
        
        # 인기 검색 종목 리스트 아이템 선택
        items = driver.find_elements(By.CSS_SELECTOR, "#popularItemList li")
        
        popular_stocks = []
        for i, item in enumerate(items[:10], 1):
            name = item.find_element(By.TAG_NAME, "a").text
            # 등락 정보 (상승, 하락, 보합 텍스트 또는 img의 alt 속성)
            # 네이버 증권은 보통 em 클래스나 img alt에 '상승', '하락' 등을 표기함
            try:
                change = item.find_element(By.TAG_NAME, "img").get_attribute("alt")
            except:
                change = "-" # 등락 정보가 없을 경우 대비
            
            price = item.find_element(By.TAG_NAME, "span").text
            
            # 형식: "1.기업명 [등락] 가격"
            popular_stocks.append(f"{i}.{name} [{change}] {price}")
        
        print("수집 완료:", popular_stocks)
        driver.quit()

        # --- 2. template.docx 생성 (수집 데이터 기반) ---
        temp_doc = Document()
        p_title = temp_doc.add_paragraph("{title}")
        set_hangul_font(p_title)
        
        temp_doc.add_paragraph("\n[실시간 수집 데이터 (등락 포함)]")
        for stock in popular_stocks:
            p = temp_doc.add_paragraph(stock)
            set_hangul_font(p)
        
        temp_doc.save('template.docx')
        print("1. 등락 정보가 포함된 template.docx 생성 완료.")

        # --- 3. template.docx를 이용한 result.docx 제작 ---
        doc = Document('template.docx')
        
        # {title} 치환
        title_text = "오늘의 인기 검색 종목 10개 입니다."
        for paragraph in doc.paragraphs:
            if '{title}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{title}', title_text)
                set_hangul_font(paragraph)

        # 하단 리스트 추가
        doc.add_paragraph("\n[최종 정리 리스트]")
        for stock in popular_stocks:
            p = doc.add_paragraph(stock, style='List Number')
            set_hangul_font(p)

        # 최종 저장
        doc.save('result.docx')
        print("2. 모든 작업이 완료된 result.docx 저장 완료.")

    except PermissionError:
        print("오류: result.docx 또는 template.docx 파일이 워드에서 열려 있습니다. 파일을 닫고 다시 실행하세요.")
    except Exception as e:
        print(f"오류 발생: {e}")

if __name__ == "__main__":
    main()