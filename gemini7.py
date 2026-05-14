import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.oxml.ns import qn

def set_hangul_font(paragraph, font_name="Malgun Gothic"):
    """한글 폰트 설정 (깨짐 방지)"""
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def main():
    chrome_options = Options()
    # 브라우저 실행 과정을 확인하기 위해 headless 모드는 사용하지 않습니다.
    
    try:
        # 1. 브라우저 실행 및 네이버 증권 접속
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=chrome_options)
        
        driver.get("https://finance.naver.com/sise/")
        time.sleep(5) # 데이터가 로드될 때까지 충분히 대기
        
        # 2. 인기 검색 종목 10개의 순수 HTML(outerHTML) 수집
        items = driver.find_elements(By.CSS_SELECTOR, "#popularItemList li")
        
        # 가공 없이 태그 전체를 리스트에 저장
        raw_html_list = []
        for item in items[:10]:
            raw_html_list.append(item.get_attribute('outerHTML'))
        
        driver.quit()

        # 3. template.docx 생성 (추가 텍스트 없이 데이터만 입력)
        doc = Document()
        
        # 첫 줄에 치환용 식별자 {title} 추가
        p_title = doc.add_paragraph("{title}")
        set_hangul_font(p_title)

        # 수집된 HTML 코드들을 곧바로 문서에 추가
        for html_code in raw_html_list:
            p = doc.add_paragraph(html_code)
            set_hangul_font(p)

        # 4. 파일 저장
        doc.save('template.docx')
        print("오류 없이 'template.docx' 파일이 생성되었습니다.")

    except Exception as e:
        print(f"오류 발생: {e}")

if __name__ == "__main__":
    main()