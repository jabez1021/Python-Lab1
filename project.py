from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.oxml.ns import qn

def set_hangul_font(paragraph, font_name="Malgun Gothic"):
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)

popular_stocks = []

try:
    driver.get('https://finance.naver.com/sise/')
    item = driver.find_element(By.ID, 'popularItemList')
    childs = item.find_elements(By.TAG_NAME, 'li')
    
    for child in childs[:10]:
        popular_stocks.append(child.text)
        print(child.text)

    try:
        document = Document('template.docx')
    except:
        document = Document()
        document.add_paragraph("{title}")

    paras = document.paragraphs
    paras[0].text = paras[0].text.replace('{title}','(제목: 오늘의 인기 검색 종목 10 개 입니다.)')
    set_hangul_font(paras[0])

    for stock in popular_stocks:
        p = document.add_paragraph(stock, style='List Number')
        set_hangul_font(p)

    document.save('result.docx')
    print("File Saved")

except Exception as e:
    print(e)

finally:
    driver.quit()