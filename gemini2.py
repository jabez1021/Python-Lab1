import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.common.exceptions import WebDriverException
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.oxml.ns import qn  # 한글 폰트 설정을 위해 필요

def set_hangul_font(paragraph, font_name="Malgun Gothic"):
    """문단의 모든 텍스트 영역에 한글 폰트를 적용하는 함수"""
    for run in paragraph.runs:
        run.font.name = font_name
        # 한글(eastAsia) 폰트 설정을 별도로 해주어야 깨지지 않습니다.
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def main():
    chrome_options = Options()
    # chrome_options.add_argument("--headless") # 필요 시 활성화
    
    try:
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=chrome_options)
    except Exception as e:
        print(f"브라우저 실행 오류: {e}")
        return

    try:
        # 1. 네이버 증권 접속 및 데이터 수집
        driver.get("https://finance.naver.com/sise/")
        time.sleep(2)
        items = driver.find_elements(By.CSS_SELECTOR, "#popularItemList li a")
        popular_stocks = [item.text for item in items[:10]]
        print(f"수집 완료: {popular_stocks}")

        # 2. 문서 로드 (template.docx)
        try:
            doc = Document('template.docx')
        except:
            doc = Document()
            doc.add_paragraph("{title}")

        # 3. 제목 변경 및 폰트 설정
        title_text = "오늘의 인기 검색 종목 10개 입니다."
        for paragraph in doc.paragraphs:
            if '{title}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{title}', title_text)
                set_hangul_font(paragraph)

        # 4. 종목 리스트 추가 및 폰트 설정
        for stock in popular_stocks:
            p = doc.add_paragraph(stock, style='List Number')
            set_hangul_font(p)

        # 5. 저장
        doc.save('result.docx')
        print("한글 설정이 적용된 result.docx 파일이 생성되었습니다.")

    finally:
        driver.quit()

if __name__ == "__main__":
    main()