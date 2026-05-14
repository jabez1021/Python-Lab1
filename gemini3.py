import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.oxml.ns import qn

def set_hangul_font(paragraph, font_name="Malgun Gothic"):
    """한글 깨짐 방지 폰트 설정"""
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def main():
    # --- Selenium 크롤링 영역 ---
    options = Options()
    try:
        driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)
        driver.get("https://finance.naver.com/sise/")
        time.sleep(2)
        
        # 인기 검색 종목 수집 (조건 3)
        items = driver.find_elements(By.CSS_SELECTOR, "#popularItemList li a")
        popular_stocks = [item.text for item in items[:10]]
        print(f"수집된 종목: {popular_stocks}")
        driver.quit()

        # --- 워드 문서 작업 영역 ---
        # 조건 4: 전달된 template.docx 파일을 템플릿으로 함 (반드시 존재해야 함)
        doc = Document('template.docx') 
        
        # {title} 부분에 제목 추가
        title_text = "오늘의 인기 검색 종목 10개 입니다."
        for paragraph in doc.paragraphs:
            if '{title}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{title}', title_text)
                set_hangul_font(paragraph) # 한글 설정

        # 조건 5: 문서 하단에 List Number 스타일로 10개 종목 추가
        for stock in popular_stocks:
            p = doc.add_paragraph(stock, style='List Number')
            set_hangul_font(p) # 한글 설정

        # 조건 6: result.docx로 저장
        doc.save('result.docx')
        print("작업이 완료되었습니다. result.docx를 확인하세요.")

    except FileNotFoundError:
        print("오류: 'template.docx' 파일이 현재 폴더에 없습니다. 파일을 준비해 주세요.")
    except Exception as e:
        print(f"오류 발생: {e}")

if __name__ == "__main__":
    main()