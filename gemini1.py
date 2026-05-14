import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.common.exceptions import WebDriverException
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document

def main():
    # --- 1. Selenium 설정 및 브라우저 실행 (예외처리 포함) ---
    chrome_options = Options()
    # chrome_options.add_argument("--headless")  # 필요 시 주석 해제 (화면 없이 실행)
    
    try:
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=chrome_options)
    except WebDriverException as e:
        print(f"브라우저 실행 중 오류가 발생했습니다: {e}")
        return

    try:
        # --- 2. 네이버 증권 국내 증시 페이지 이동 ---
        url = "https://finance.naver.com/sise/"
        driver.get(url)
        time.sleep(2)  # 페이지 로딩 대기

        # --- 3. 인기 검색 종목명 수집 및 출력 ---
        print("인기 검색 종목을 수집 중입니다...")
        # 네이버 증권 '인기 검색 종목' 영역의 셀렉터
        items = driver.find_elements(By.CSS_SELECTOR, "#popularItemList li a")
        
        popular_stocks = []
        for item in items[:10]:  # 최대 10개 수집
            popular_stocks.append(item.text)
        
        print(f"수집된 종목: {popular_stocks}")

        # --- 4. template.docx 로드 및 제목 수정 ---
        # 주의: 코드 실행 경로에 template.docx 파일이 있어야 합니다.
        try:
            doc = Document('template.docx')
        except Exception:
            print("template.docx 파일을 찾을 수 없어 새로운 문서를 생성합니다.")
            doc = Document()

        # {title} 부분을 찾아 교체
        title_text = "오늘의 인기 검색 종목 10개 입니다."
        for paragraph in doc.paragraphs:
            if '{title}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{title}', title_text)

        # --- 5. 문서 하단에 List Number 스타일로 종목 추가 ---
        for stock in popular_stocks:
            doc.add_paragraph(stock, style='List Number')

        # --- 6. result.docx로 저장 ---
        doc.save('result.docx')
        print("성공적으로 result.docx 파일이 저장되었습니다.")

    except Exception as e:
        print(f"작업 중 예상치 못한 오류가 발생했습니다: {e}")
    
    finally:
        driver.quit()

if __name__ == "__main__":
    main()